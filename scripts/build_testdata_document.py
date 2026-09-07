"""Build the test users and test data reference document.

Two files are produced:

  reports/test-data-reference.docx   - the dataset. Safe to share and to
                                       include in the submission. No passwords.

  .local/TEST-USERS-AND-PASSWORDS.docx - the same thing plus the sign-in
                                       details. Written to .local/, which is
                                       git-ignored, because the repository is
                                       public.

Run: python scripts/build_testdata_document.py
"""

from __future__ import annotations

import json
import sys
from datetime import date, timedelta
from pathlib import Path

# The script lives in scripts/, so the project root has to be importable
# before app.testdata can be read.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.testdata import (
    BUSY_HOLDER,
    BUSY_HOLDER_CREDENTIALS,
    DEMO_CREDENTIALS,
    QUALIFICATIONS,
    UNIVERSITIES,
    USER_NOTES,
    USERS,
    VERIFICATIONS,
)

ROOT = Path(__file__).resolve().parent.parent
PUBLIC_OUT = ROOT / "reports" / "test-data-reference.docx"
PRIVATE_DIR = ROOT / ".local"
PRIVATE_OUT = PRIVATE_DIR / "TEST-USERS-AND-PASSWORDS.docx"
CREDENTIALS_FILE = PRIVATE_DIR / "live-test-credentials.json"

LIVE_URL = "https://qvs-mim736.fly.dev"
REPO_URL = "https://github.com/SimonMutyavaviri/qualification-verification-system"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x47, 0x55, 0x69)
RED = RGBColor(0x9F, 0x12, 0x39)
GREEN = RGBColor(0x06, 0x5F, 0x46)

TODAY = date.today()


def _on(days_ago: int) -> str:
    return (TODAY - timedelta(days=days_ago)).strftime("%d %b %Y")


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------


def shade(paragraph, hex_fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    properties.append(shading)


def mono(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    shade(p, "F1F5F9")


def callout(
    doc: Document, title: str, text: str, colour: RGBColor = SLATE, fill: str = "FFFBEB"
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    t = p.add_run(f"{title}  ")
    t.bold = True
    t.font.color.rgb = colour
    t.font.size = Pt(10.5)
    b = p.add_run(text)
    b.font.size = Pt(10.5)
    b.font.color.rgb = NAVY
    shade(p, fill)


def h(doc: Document, text: str, level: int = 1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = NAVY
    return heading


def para(doc: Document, text: str, size: float = 11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    font: float = 9.5,
) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = head
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12
    for level, size in ((1, 17), (2, 13.5), (3, 12)):
        s = doc.styles[f"Heading {level}"]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = NAVY


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------


def title_page(doc: Document, private: bool) -> None:
    t = doc.add_heading("Test Users and Test Data", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in t.runs:
        r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Qualification Verification System\nMIM736 Practical Assignment")
    r.font.size = Pt(13)
    r.font.color.rgb = SLATE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Live system: {LIVE_URL}")
    r.font.size = Pt(11)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Data loaded: {TODAY.strftime('%d %B %Y')}")
    r.font.size = Pt(10)
    r.font.color.rgb = SLATE

    doc.add_paragraph()
    if private:
        callout(
            doc,
            "This copy contains passwords. Do not commit or submit it.",
            "It is stored in the .local folder, which Git ignores, because the "
            "repository is public. A version without passwords, safe to submit, "
            "is at reports/test-data-reference.docx.",
            RED,
            "FFF1F2",
        )
    else:
        callout(
            doc,
            "This copy contains no passwords.",
            "It is safe to include in the submission and to share. The sign-in "
            "details are held separately by the project owner.",
            GREEN,
            "ECFDF5",
        )
    doc.add_page_break()


def overview(doc: Document) -> None:
    h(doc, "1. What is loaded, and why")

    para(
        doc,
        "The live system holds a test register built around three "
        "universities. The data is not random: it was chosen so that every "
        "screen, every verification result and every review-assistant rule can "
        "be demonstrated with a real record rather than described in words.",
    )

    table(
        doc,
        ["What", "How many"],
        [
            ["Universities (active)", "3"],
            ["Staff accounts", str(len(USERS))],
            ["Qualifications in the register", "40"],
            ["Verification attempts recorded", "40"],
            ["Audit entries", "44 and rising"],
        ],
        [3.2, 1.4],
        10.5,
    )

    h(doc, "The three universities", level=2)
    table(
        doc,
        ["Code", "Name", "Country", "Registry contact"],
        [[u["code"], u["name"], u["country"], u["contact_email"]] for u in UNIVERSITIES],
        [0.7, 2.6, 1.0, 1.9],
    )

    callout(
        doc,
        "A fourth institution appears, deactivated.",
        "An earlier professional body, ICA, is still listed but marked "
        "Inactive. It was deactivated rather than deleted, so no new "
        "credentials can be registered against it while the ones it already "
        "issued stay verifiable. That behaviour is worth demonstrating.",
    )

    doc.add_page_break()


def users_section(doc: Document, private: bool, creds: dict | None) -> None:
    h(doc, "2. Test users")

    para(
        doc,
        "Seven accounts, covering all three roles. Sign in at "
        f"{LIVE_URL}/login using either the username or the email address.",
    )

    role_label = {
        "admin": "Administrator",
        "issuer": "Qualification Issuer",
        "verifier": "Verifier",
    }

    if private and creds:
        rows = []
        for username, full_name, role, code, _env in USERS:
            rows.append(
                [
                    username,
                    creds.get(username, "[not recorded]"),
                    full_name,
                    role_label[role.value],
                    code or "-",
                ]
            )
        table(
            doc,
            ["Username", "Password", "Name", "Role", "University"],
            rows,
            [1.15, 1.35, 1.5, 1.35, 0.85],
        )
        callout(
            doc,
            "Passwords are shown once at seeding and cannot be recovered.",
            "If one is lost, sign in as an administrator and set a new one from "
            "the user's profile page, or create a replacement account.",
        )
    else:
        rows = []
        for username, full_name, role, code, _env in USERS:
            rows.append([username, full_name, role_label[role.value], code or "-"])
        table(
            doc,
            ["Username", "Name", "Role", "University"],
            rows,
            [1.4, 1.9, 1.7, 1.2],
        )
        callout(
            doc,
            "Passwords are deliberately not printed here.",
            "The repository is public. The project owner holds them separately.",
        )

    h(doc, "What each account is for", level=2)
    table(
        doc,
        ["Username", "Use it to show"],
        [[u, USER_NOTES[u]] for u, _n, _r, _c, _e in USERS],
        [1.4, 4.8],
    )

    h(doc, "What each role can do", level=2)
    table(
        doc,
        ["Action", "Verifier", "Issuer", "Admin"],
        [
            ["Verify a credential", "Yes", "Yes", "Yes"],
            ["Search the register", "Yes", "Yes", "Yes"],
            ["See own verification history", "Yes", "Yes", "Yes"],
            ["See everyone's verification history", "No", "Yes", "Yes"],
            ["Register a qualification", "No", "Yes", "Yes"],
            ["Edit a qualification", "No", "Yes", "Yes"],
            ["Revoke a qualification", "No", "Yes", "Yes"],
            ["Reinstate a revoked qualification", "No", "No", "Yes"],
            ["See the review assistant", "No", "Yes", "Yes"],
            ["Manage users and institutions", "No", "No", "Yes"],
            ["View the audit trail", "No", "No", "Yes"],
        ],
        [3.0, 1.05, 1.05, 1.05],
    )

    callout(
        doc,
        "Worth demonstrating.",
        "Sign in as hr.verifier and type /admin/users into the address bar. You "
        "get 403 Access denied, and the attempt is written to the audit trail. "
        "Hiding a menu is not access control; this is.",
        GREEN,
        "ECFDF5",
    )

    doc.add_page_break()


def demo_section(doc: Document) -> None:
    h(doc, "3. Credentials to use in the demonstration")

    para(
        doc,
        "These ten cover every outcome the system can produce. Type them into "
        "the Verify page exactly as written; capitalisation does not matter.",
    )

    tone = {
        "VALID": "Green tick",
        "EXPIRED": "Amber warning",
        "REVOKED": "Red cross",
        "INVALID": "Red cross",
    }
    table(
        doc,
        ["Credential ID", "Result", "Shows as", "Why it is worth showing"],
        [[cid, res.value, tone[res.value], note] for cid, res, note in DEMO_CREDENTIALS],
        [1.5, 0.85, 1.05, 2.8],
    )

    callout(
        doc,
        "The most important pair.",
        "Show MSU-2019-CER274 (EXPIRED) next to NUST-2020-BEN538 (REVOKED). "
        "Both are 'not valid', but one is a genuine award that lapsed and the "
        "other was withdrawn as fraudulent. A system that returned only "
        "'invalid' would hide that difference from the employer who most needs "
        "it.",
        GREEN,
        "ECFDF5",
    )

    h(doc, "Showing the review assistant", level=2)
    para(
        doc,
        "Sign in as an issuer or administrator, open a qualification from the "
        "register, and look at the panel on the right. Each record below trips "
        "a different rule.",
    )
    table(
        doc,
        ["Open this record", "Rule it triggers", "What the panel says"],
        [
            [
                "MSU-2020-BSC352",
                "incomplete_holder_contact",
                "No holder email recorded, so the award cannot be confirmed with " "the holder",
            ],
            [
                "UZ-1955-BSC486",
                "implausible_award_date",
                "Award date is about 70 years ago, beyond the plausibility window",
            ],
            [
                "NUST-2025-SHC672",
                "short_validity_window",
                "Valid for only 14 days, which is unusually short",
            ],
            [
                "MSU-2019-SHC423",
                "high_holder_volume",
                "Six credentials are registered to this holder name",
            ],
            [
                "NUST-2020-BEN538",
                "repeated_failed_lookups",
                "Five verification attempts on this reference did not return VALID",
            ],
            ["UZ-2021-DIP658", "revoked_credential", "This credential has been revoked"],
        ],
        [1.4, 1.75, 3.05],
    )

    callout(
        doc,
        "Say this in the viva.",
        "The assistant is rules, not machine learning, on purpose. Something "
        "that influences whether a person's degree is believed has to be able "
        "to explain itself, and every signal here points at a named field and "
        "gives a reason.",
    )

    doc.add_page_break()


def register_section(doc: Document) -> None:
    h(doc, "4. The full register")

    para(
        doc,
        "Every qualification loaded, grouped by university. 'Status' is what "
        "the record says; 'Verifies as' is what a verifier actually sees, "
        "which differs where a credential has passed its expiry date.",
    )

    rows_by_code: dict[str, list[list[str]]] = {"MSU": [], "UZ": [], "NUST": []}

    all_rows = list(QUALIFICATIONS)
    for cid, title, qtype, code, award_ago in BUSY_HOLDER_CREDENTIALS:
        all_rows.append(
            (
                cid,
                title,
                qtype,
                BUSY_HOLDER,
                "tapiwa.chidziva@example.com",
                code,
                award_ago,
                None,
                None,
                None,
            )
        )

    for entry in all_rows:
        cid, title, qtype, holder, email, code, award_ago, expiry_ago = entry[:8]
        status = entry[8]
        status_name = status.value.capitalize() if status is not None else "Active"

        if status is not None and status.value == "revoked":
            verifies = "REVOKED"
        elif expiry_ago is not None and expiry_ago > 0:
            verifies = "EXPIRED"
        else:
            verifies = "VALID"

        expiry = "-" if expiry_ago is None else _on(expiry_ago)
        rows_by_code[code].append(
            [cid, title, qtype, holder, _on(award_ago), expiry, status_name, verifies]
        )

    names = {u["code"]: u["name"] for u in UNIVERSITIES}
    for code in ("MSU", "UZ", "NUST"):
        h(doc, f"{names[code]} ({code}) - {len(rows_by_code[code])} credentials", level=2)
        table(
            doc,
            [
                "Credential ID",
                "Qualification",
                "Type",
                "Holder",
                "Awarded",
                "Expires",
                "Status",
                "Verifies as",
            ],
            rows_by_code[code],
            [1.15, 1.55, 0.85, 1.05, 0.72, 0.72, 0.6, 0.66],
            8,
        )

    callout(
        doc,
        "Three older records also exist.",
        "QVS-2023-DEMO01, QVS-2022-DEMO02 and QVS-2021-DEMO03 remain from the "
        "first deployment and still verify as valid, expired and revoked. "
        "QVS-2022-DEMO02 belongs to the deactivated institution, which is a "
        "neat way to show that deactivating an institution does not invalidate "
        "the credentials it already issued.",
    )

    doc.add_page_break()


def history_section(doc: Document) -> None:
    h(doc, "5. Verification history")

    para(
        doc,
        f"{len(VERIFICATIONS)} verification attempts were recorded, so the "
        "history page, the dashboard counters and the audit trail all have "
        "real content. The spread is deliberate.",
    )

    counts: dict[str, int] = {}
    for _cid, user, _ago in VERIFICATIONS:
        counts[user] = counts.get(user, 0) + 1
    table(
        doc,
        ["Performed by", "Attempts"],
        [[u, str(n)] for u, n in sorted(counts.items())],
        [2.4, 1.2],
        10.5,
    )

    h(doc, "References that were never registered", level=2)
    para(
        doc,
        "These were checked deliberately so the register shows what a failed " "lookup looks like:",
    )
    table(
        doc,
        ["Reference tried", "Times", "Point being made"],
        [
            ["MSU-2023-FAKE47", "1", "A single failed check"],
            ["UZ-2024-BOGUS3", "1", "A failed check by a different verifier"],
            ["NUST-2022-NONE24", "1", "A third, at a third university"],
            [
                "MSU-2022-TRY777",
                "5",
                "The same unregistered reference tried five times. This is the "
                "pattern a fraud investigation would look for, and it is only "
                "visible because failed attempts are recorded.",
            ],
        ],
        [1.5, 0.7, 4.0],
    )

    callout(
        doc,
        "Why failed attempts are stored at all.",
        "Most systems record only successful lookups. Ours stores every "
        "attempt, including malformed ones, because 'somebody tried to verify "
        "a credential that does not exist, from this address, at this time' is "
        "exactly the signal that matters.",
        GREEN,
        "ECFDF5",
    )

    doc.add_page_break()


def howto_section(doc: Document) -> None:
    h(doc, "6. Reloading or extending the data")

    para(doc, "The dataset is code, not a database dump. It lives in:")
    mono(doc, "app/testdata.py")

    para(doc, "To load it into a local copy of the system:")
    mono(doc, "set FLASK_APP=run.py\n" "flask seed-testdata")

    para(
        doc,
        "The command is safe to run more than once. It skips anything already "
        "present, so it will not create duplicates.",
    )

    para(doc, "To start again from an empty register:")
    mono(doc, "flask seed-testdata --reset")

    callout(
        doc,
        "--reset deletes qualifications, verifications and audit entries.",
        "Do not run it against the live system close to the demonstration "
        "unless you intend to lose the history.",
        RED,
        "FFF1F2",
    )

    para(doc, "To choose the passwords rather than have them generated:")
    mono(
        doc,
        "set SEED_ADMIN_PASSWORD=YourPassword123\n"
        "set SEED_MSU_PASSWORD=AnotherPassword123\n"
        "flask seed-testdata",
    )

    para(doc, "To load it into the live system on Fly.io:")
    mono(
        doc,
        "flyctl ssh console --app qvs-mim736 \\\n"
        '  -C "gosu qvs env FLASK_APP=wsgi.py python -m flask seed-testdata"',
    )

    h(doc, "Adding your own records", level=2)
    para(doc, "Open app/testdata.py and add a row to the QUALIFICATIONS list. Each " "row is:")
    mono(
        doc,
        "(credential_id, title, type, holder, holder_email, university_code,\n"
        " award_days_ago, expiry_days_ago_or_None, status, revocation_reason)",
    )
    para(
        doc,
        "A negative expiry means a date in the future, so the credential is "
        "still valid. Credential IDs must match PREFIX-YEAR-SIXCHARS and "
        "should avoid the characters O, 0, I and 1, which get misread off a "
        "printed certificate.",
    )

    doc.add_page_break()


def checklist_section(doc: Document) -> None:
    h(doc, "7. Demonstration checklist")

    para(doc, "A run through the data that covers everything the marking criteria " "ask about.")

    table(
        doc,
        ["#", "Do this", "Sign in as", "Done"],
        [
            ["1", "Sign in and show the dashboard counters moving", "admin", ""],
            ["2", "Search the register for 'Tinashe'", "admin", ""],
            ["3", "Filter the register by status = Revoked (4 records)", "admin", ""],
            ["4", "Verify MSU-2023-BSC247 - VALID", "hr.verifier", ""],
            ["5", "Verify MSU-2019-CER274 - EXPIRED", "hr.verifier", ""],
            ["6", "Verify NUST-2020-BEN538 - REVOKED", "hr.verifier", ""],
            ["7", "Verify MSU-2023-FAKE47 - INVALID, nothing disclosed", "hr.verifier", ""],
            ["8", "Open the verification receipt from the result page", "hr.verifier", ""],
            ["9", "Show that a verifier sees only their own history", "hr.verifier", ""],
            ["10", "Register a brand new qualification", "registry.msu", ""],
            ["11", "Verify the one you just registered - VALID", "registry.msu", ""],
            ["12", "Revoke it with a reason, then verify again - REVOKED", "registry.msu", ""],
            ["13", "Open NUST-2020-BEN538 and show the review assistant", "registry.nust", ""],
            ["14", "Type /admin/users in the address bar - 403 Access denied", "hr.verifier", ""],
            ["15", "Show that 403 in the audit trail", "admin", ""],
            ["16", "Show the full audit trail and filter it", "admin", ""],
            ["17", "Show the three universities and the deactivated one", "admin", ""],
        ],
        [0.35, 3.5, 1.35, 0.5],
        10,
    )

    callout(
        doc,
        "Steps 10 to 12 are the strongest part.",
        "Registering a credential, verifying it, revoking it and watching the "
        "same reference change from VALID to REVOKED proves the whole system "
        "works end to end on live data. It is far more convincing than a tour "
        "of the screens.",
        GREEN,
        "ECFDF5",
    )

    h(doc, "Where to find more", level=2)
    table(
        doc,
        ["Document", "What it covers"],
        [
            ["reports/demonstration-script.md", "A word-for-word 13-minute script"],
            ["reports/viva-preparation.md", "48 likely questions with answers"],
            ["docs/user-manual.md", "How to use every screen"],
            ["docs/evidence-checklist.md", "Every screenshot to capture"],
            ["docs/test-cases.md", "The formal test cases"],
        ],
        [2.4, 3.8],
    )

    para(doc, f"Repository: {REPO_URL}")
    para(doc, f"Live system: {LIVE_URL}")


# ---------------------------------------------------------------------------


def build(private: bool, creds: dict | None, out: Path) -> None:
    doc = Document()
    style_document(doc)
    title_page(doc, private)
    overview(doc)
    users_section(doc, private, creds)
    demo_section(doc)
    register_section(doc)
    history_section(doc)
    howto_section(doc)
    checklist_section(doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"  {out.relative_to(ROOT)}")


def main() -> None:
    creds = None
    if CREDENTIALS_FILE.exists():
        creds = json.loads(CREDENTIALS_FILE.read_text(encoding="utf-8"))

    print("Building test data documents:\n")
    build(False, None, PUBLIC_OUT)
    if creds:
        build(True, creds, PRIVATE_OUT)
    else:
        print("  (no .local/live-test-credentials.json - skipped the password copy)")
    print("\nDone.")


if __name__ == "__main__":
    main()
