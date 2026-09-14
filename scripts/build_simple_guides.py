"""Build a very short, plain-language first-step guide for each teammate.

The full team guides assume more than a non-technical reader can absorb. This
one does a single thing: gets each person's name onto the project by editing
the top of their own individual report on the GitHub website, then getting it
approved and merged.

Each person edits a different file, so all three can do it at the same time
without causing a merge conflict. The pull request targets main, so the change
counts on the repository's Contributors list as soon as it is merged.

Run: python scripts/build_simple_guides.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "reports" / "team-guides"

REPO_URL = "https://github.com/SimonMutyavaviri/qualification-verification-system"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
GREEN = RGBColor(0x06, 0x5F, 0x46)
RED = RGBColor(0x9F, 0x12, 0x39)
GREY = RGBColor(0x47, 0x55, 0x69)

PEOPLE = [
    {"handle": "Learnmore1988", "student": 1, "reviewer": "AnesuCK"},
    {"handle": "Progress861", "student": 2, "reviewer": "SimonMutyavaviri"},
    {"handle": "AnesuCK", "student": 3, "reviewer": "Progress861"},
]


def shade(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY


def plain(doc: Document, text: str, size: float = 13, colour: RGBColor = NAVY) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = colour


def do(doc: Document, number: int, *parts: str) -> None:
    """One numbered action. Parts wrapped in ** are shown bold (button names)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    n = p.add_run(f"{number}.   ")
    n.bold = True
    n.font.size = Pt(14)
    n.font.color.rgb = GREEN
    for part in parts:
        bold = part.startswith("**") and part.endswith("**")
        r = p.add_run(part.strip("*"))
        r.bold = bold
        r.font.size = Pt(13)


def box(doc: Document, text: str, fill: str = "F1F5F9", mono: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    if mono:
        r.font.name = "Consolas"
    shade(p, fill)


def note(doc: Document, title: str, text: str, colour: RGBColor, fill: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    t = p.add_run(title + "  ")
    t.bold = True
    t.font.size = Pt(13)
    t.font.color.rgb = colour
    b = p.add_run(text)
    b.font.size = Pt(13)
    shade(p, fill)


def build(person: dict) -> Path:
    handle = person["handle"]
    number = person["student"]
    reviewer = person["reviewer"]
    file_path = f"reports/student-{number}-individual-report.md"
    file_url = f"{REPO_URL}/blob/main/{file_path}"

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Your first step on GitHub")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"For: {handle}      Takes about 15 minutes      Nothing to install")
    r.font.size = Pt(13)
    r.font.color.rgb = GREY

    note(
        doc,
        "What you are doing:",
        "typing your name into your own page of the project. GitHub records "
        "that YOU did it. That record is part of your mark. You cannot break "
        "anything.",
        GREEN,
        "ECFDF5",
    )

    # ------------------------------------------------------------------
    heading(doc, "Part 1 - Type your name")

    do(doc, 1, "Sign in to GitHub as ", f"**{handle}**", ".")
    do(doc, 2, "Open this link (copy it into your browser):")
    box(doc, file_url)
    do(
        doc,
        3,
        "You will see a page called ",
        "**Individual Contribution Report**",
        ". Near the top right of the page there is a small ",
        "**pencil**",
        " picture. Click it.",
    )
    plain(
        doc,
        "If you cannot see a pencil, you are not signed in. Go back to step 1.",
        12,
        GREY,
    )
    do(doc, 4, "A box opens where you can type. Find these three lines near the top:")
    box(
        doc,
        "| **Full name** | `[INSERT YOUR FULL NAME]` |\n"
        "| **Student number** | `[INSERT YOUR STUDENT NUMBER]` |\n"
        "| **GitHub username** | `[INSERT YOUR GITHUB USERNAME]` |",
    )
    do(
        doc,
        5,
        "Change them so they look like this, with your real name and number. "
        "Delete the square brackets and the little ` marks too:",
    )
    box(
        doc,
        "| **Full name** | Your Real Name |\n"
        "| **Student number** | Your student number |\n"
        f"| **GitHub username** | {handle} |",
        "ECFDF5",
    )
    note(
        doc,
        "Only change those three lines.",
        "Leave everything else exactly as it is, including the | symbols.",
        RED,
        "FFF1F2",
    )
    do(doc, 6, "Click the green button at the top right. It says ", "**Commit changes**", ".")
    do(doc, 7, "A small window opens. In the first box, type exactly:")
    box(doc, "docs: add my details to my individual report")
    do(
        doc,
        8,
        "Below that, make sure the option ",
        "**Create a new branch for this commit and start a pull request**",
        " is selected. (It usually already is.)",
    )
    do(doc, 9, "Click the green ", "**Propose changes**", " button.")
    do(
        doc,
        10,
        "A new page opens. Click the green ",
        "**Create pull request**",
        " button. If a second green ",
        "**Create pull request**",
        " button appears, click that one too.",
    )

    note(
        doc,
        "Well done.",
        "You have made your first change. Now it needs to be checked.",
        GREEN,
        "ECFDF5",
    )

    # ------------------------------------------------------------------
    heading(doc, "Part 2 - Get it checked, then finish")

    do(
        doc,
        11,
        "Stay on the page you are on. Wait about ",
        "**3 minutes**",
        ", then refresh the page (press F5).",
    )
    do(
        doc,
        12,
        "Scroll to the bottom. Wait until you see a green tick and the words ",
        "**All checks have passed**",
        ". The computer is checking your change for you.",
    )
    do(
        doc,
        13,
        "Copy the web address of this page from the top of your browser. Send it " "to ",
        f"**{reviewer}**",
        " on WhatsApp with the message: “Please approve my GitHub change.”",
    )
    do(
        doc,
        14,
        f"When {reviewer} tells you it is approved, open the same page again and "
        "click the green ",
        "**Squash and merge**",
        " button, then ",
        "**Confirm squash and merge**",
        ".",
    )
    do(doc, 15, "If it offers ", "**Delete branch**", ", click it. You are finished.")

    note(
        doc,
        "If the button says 'Update branch' instead:",
        "click Update branch, wait 3 minutes for the green tick again, then do " "step 14.",
        GREY,
        "F1F5F9",
    )

    # ------------------------------------------------------------------
    heading(doc, "Part 3 - When a teammate asks YOU to check theirs")

    do(doc, 1, "Open the link they send you. Make sure you are signed in.")
    do(doc, 2, "Click the tab ", "**Files changed**", " near the top.")
    do(doc, 3, "Check they only changed their name, number and username. Nothing else.")
    do(doc, 4, "Click the green button ", "**Review changes**", " at the top right.")
    do(doc, 5, "In the box, write one sentence, for example:")
    box(
        doc,
        "Checked - only the student details were changed and the checks passed.",
        mono=False,
    )
    do(doc, 6, "Choose ", "**Approve**", ", then click ", "**Submit review**", ".")
    do(doc, 7, "Send them a message: “Approved - you can merge now.”")

    # ------------------------------------------------------------------
    heading(doc, "Part 4 - Take 3 screenshots")

    plain(
        doc,
        "Press the Windows key + Shift + S together, drag a box over the screen, "
        "then paste it into a Word document with Ctrl + V. Keep the web address "
        "visible at the top.",
        13,
    )
    do(doc, 1, "Your change, on the ", "**Files changed**", " tab.")
    do(doc, 2, "The green ", "**All checks have passed**", " message.")
    do(doc, 3, "The purple ", "**Merged**", " label after you finish.")

    # ------------------------------------------------------------------
    heading(doc, "Stuck?")

    plain(
        doc,
        "Take a screenshot of what you see and send it to SimonMutyavaviri. "
        "Do not worry - nothing you do here can damage the project.",
        13,
    )

    path = OUT / f"START-HERE-{handle}.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for person in PEOPLE:
        print(f"  {build(person).relative_to(ROOT)}")


if __name__ == "__main__":
    main()
